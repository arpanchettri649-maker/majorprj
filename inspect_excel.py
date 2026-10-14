import pandas as pd
import openpyxl
from openpyxl import load_workbook

print("=" * 80)
print("INSPECTING: test_attendance_2024-2027.xlsx")
print("=" * 80)

try:
    wb = load_workbook("test_attendance_2024-2027.xlsx")
    print(f"Sheet names: {wb.sheetnames}")
    
    for sheet_name in wb.sheetnames:
        print(f"\n--- Sheet: {sheet_name} ---")
        ws = wb[sheet_name]
        print(f"Dimensions: {ws.dimensions}")
        print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
        
        # Print first 15 rows
        print("\nFirst 15 rows:")
        for row_idx in range(1, min(16, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(11, ws.max_column + 1)):
                cell = ws.cell(row=row_idx, column=col_idx)
                row_data.append(str(cell.value)[:30] if cell.value else "")
            print(f"Row {row_idx}: {row_data}")
        
        # Check for merged cells
        print(f"\nMerged cells in {sheet_name}:")
        for merge in ws.merged_cells.ranges:
            print(f"  {merge}")
    
    wb.close()
except Exception as e:
    print(f"Error reading with openpyxl: {e}")

# Try with pandas
try:
    print("\n" + "=" * 80)
    print("READING WITH PANDAS (no header)")
    print("=" * 80)
    df = pd.read_excel("test_attendance_2024-2027.xlsx", header=None, nrows=20)
    print(f"Shape: {df.shape}")
    print("\nFirst 20 rows:")
    print(df.to_string())
except Exception as e:
    print(f"Error reading with pandas: {e}")

print("\n" + "=" * 80)
print("INSPECTING: test_attendance_dce5.xlsx")
print("=" * 80)

try:
    wb = load_workbook("test_attendance_dce5.xlsx")
    print(f"Sheet names: {wb.sheetnames}")
    
    for sheet_name in wb.sheetnames:
        print(f"\n--- Sheet: {sheet_name} ---")
        ws = wb[sheet_name]
        print(f"Dimensions: {ws.dimensions}")
        print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")
        
        # Print first 15 rows
        print("\nFirst 15 rows:")
        for row_idx in range(1, min(16, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, min(11, ws.max_column + 1)):
                cell = ws.cell(row=row_idx, column=col_idx)
                row_data.append(str(cell.value)[:30] if cell.value else "")
            print(f"Row {row_idx}: {row_data}")
        
        # Check for merged cells
        print(f"\nMerged cells in {sheet_name}:")
        for merge in ws.merged_cells.ranges:
            print(f"  {merge}")
    
    wb.close()
except Exception as e:
    print(f"Error reading with openpyxl: {e}")

# Try with pandas
try:
    print("\n" + "=" * 80)
    print("READING WITH PANDAS (no header)")
    print("=" * 80)
    df = pd.read_excel("test_attendance_dce5.xlsx", header=None, nrows=20)
    print(f"Shape: {df.shape}")
    print("\nFirst 20 rows:")
    print(df.to_string())
except Exception as e:
    print(f"Error reading with pandas: {e}")
